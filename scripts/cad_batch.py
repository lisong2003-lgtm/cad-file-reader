#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
cad_batch.py — 按目录批量扫图，输出识图台账（Excel + 自检汇总.md）。

设计：
  * 复用现有 scripts/cad_scan.py 作为单图解析引擎（通过 run_cad.sh 封装调用），不新写整套解析。
  * 每张图独立运行；单图失败只记入台账，不中断整批。
  * 每张图用 cad_contract.contractize_payload 包成 cad-file-reader/v0 契约
    （逐候选 final_quantity=false、source_schema=cad-file-reader/v0），并用现有校验器核验。
  * --resume 按源文件 SHA-256 断点续跑（state 文件 <out>/_state.json）。
  * 输出：<out>/识图台账.xlsx（黄底标出需人工复核行）、<out>/自检汇总.md、
    <out>/_contract/*.json（v0 契约）、<out>/_detail/*.json（cad_scan 详情）、<out>/_json/*.json（单图 cad_scan JSON）。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path
from typing import Any

from cad_contract import contractize_payload, validate_contract_payload
from cad_preflight import input_fail_reason, preflight_file

SKILL_DIR = Path(__file__).resolve().parents[1]
RUN_CAD_SH = SKILL_DIR / "scripts" / "run_cad.sh"
CONTRACT_SCHEMA = "cad-file-reader/v0"
TINY_MIN_BYTES = 1024  # <1KB 视为缺输入（快速失败）

YELLOW = "FFF2CC"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_stem(path: Path) -> str:
    name = re.sub(r"[^A-Za-z0-9_\u4e00-\u9fff-]+", "_", path.stem).strip("_. ")
    return name or path.name


def classify_input(path: Path) -> str:
    """返回 'ok' | '缺输入' | '失败跳过'（快速失败预检，不做解析）。"""
    pf = preflight_file(path)
    if pf.get("ok"):
        return "ok"
    if pf.get("kind") in ("not_found", "not_file", "not_readable", "empty", "tiny"):
        return "缺输入"
    return "失败跳过"


def magic_reason(path: Path) -> str:
    return input_fail_reason(path) or "魔数/头部不是标准 DWG 或 DXF"


def collect_files(paths: list[str], glob_pat: str, recursive: bool) -> list[Path]:
    import fnmatch
    files: list[Path] = []
    for p in paths:
        pp = Path(p).expanduser()
        if pp.is_dir():
            it = pp.rglob("*") if recursive else pp.glob("*")
            files += sorted(
                x for x in it
                if x.is_file() and fnmatch.fnmatch(x.name.lower(), glob_pat.lower())
            )
        elif pp.is_file():
            files.append(pp)
    out: list[Path] = []
    seen: set[str] = set()
    for f in files:
        key = str(f.resolve())
        if key not in seen:
            seen.add(key)
            out.append(f)
    return out


def run_cad_scan(path: Path, detail_path: Path, json_path: Path, budget: float, cache_dir: str | None) -> tuple[int, str]:
    cmd: list[str] = []
    if RUN_CAD_SH.exists():
        # 走 bash 启动：SkillHub/ClawHub 打包会丢失 .sh 执行位
        cmd = ["bash", str(RUN_CAD_SH), "cad_scan"]
    else:
        # 退化：直接调 python（依赖 PYTHONPATH 已含 scripts 同级路径）
        cmd = [sys.executable, str(SKILL_DIR / "scripts/cad_scan.py")]
    cmd += [str(path), "--format", "json", "--detail-json", str(detail_path), "--budget", str(budget)]
    if cache_dir:
        cmd += ["--cache-dir", str(cache_dir)]
    env = dict(os.environ)
    env.setdefault("CAD_CONCISE", "1")
    timeout = budget if budget and budget > 0 else None
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        return 124, (exc.stderr or "")[-3000:] + f"；cad_scan 超时终止（预算 {budget}s）"
    except Exception as exc:  # noqa: BLE001
        return 126, f"cad_scan 启动失败：{exc}"
    detail_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    if proc.stdout and proc.stdout.strip():
        try:
            json_path.write_text(proc.stdout, encoding="utf-8")
        except OSError:
            pass
    if proc.returncode != 0:
        err = (proc.stderr or "")[-3000:]
        return proc.returncode, err or "cad_scan 退出非零"
    return 0, ""


def build_contract(detail: dict[str, Any], source_name: str, source_path: Path, digest: str) -> dict[str, Any]:
    """把 cad_scan 详情 JSON 里的 text_records 包装成 cad-file-reader/v0 契约。"""
    text_rows: list[dict[str, Any]] = []
    for i, rec in enumerate(detail.get("text_records") or []):
        text = str(rec.get("text") or "").strip()[:240]
        layer = str(rec.get("layer") or "")
        x = rec.get("x")
        y = rec.get("y")
        bbox: list[float] | None = None
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            pad = 5.0
            bbox = [max(0.0, float(x) - pad), max(0.0, float(y) - pad),
                    float(x) + pad, float(y) + pad]
        confidence = 0.85 if (text and layer and bbox) else 0.40
        row: dict[str, Any] = {
            "id": f"txt{i + 1}",
            "text": text,
            "layer": layer,
            "bbox": bbox,
            "confidence": round(confidence, 4),
            "status": "candidate" if confidence >= 0.50 else "review",
            "method": "cad_scan 低内存文字提取",
        }
        if confidence < 0.50:
            row["review_reason"] = "低置信：缺图层或证据坐标不足，需人工复核"
        text_rows.append(row)
    payload: dict[str, Any] = {
        "schema": "cad-scan-detail",
        "source_file": source_name,
        "source_path": str(source_path),
        "input_sha256": digest,
        "text_records": text_rows,
        "text": text_rows[:0],
    }
    contractize_payload(payload, source=source_name, concise=False)
    contract = payload.get("contract") or {}
    contract.setdefault("source_schema", CONTRACT_SCHEMA)
    payload["contract"] = contract
    return payload


def coord_of(row: dict[str, Any]) -> str:
    bbox = row.get("bbox")
    if isinstance(bbox, (list, tuple)) and len(bbox) >= 4:
        try:
            x1, y1, x2, y2 = (float(v) for v in bbox[:4])
            return f"{x1:.0f},{y1:.0f}~{x2:.0f},{y2:.0f}"
        except (TypeError, ValueError):
            pass
    ev = row.get("evidence")
    if isinstance(ev, list):
        for item in ev:
            if isinstance(item, dict):
                x = item.get("x")
                y = item.get("y")
                if isinstance(x, (int, float)) and isinstance(y, (int, float)):
                    return f"{float(x):.0f},{float(y):.0f}"
    return ""


def review_reason_text(row: dict[str, Any]) -> str:
    parts: list[str] = []
    notes = str(row.get("review_notes") or "").strip()
    if notes:
        parts.append(notes)
    for r in row.get("review_reasons") or []:
        if r not in parts:
            parts.append(r)
    return "；".join(parts)


def flatten_candidates(digest: str, source_name: str, contract: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for arr_name, tier_label in (("candidates", "候选"), ("review_candidates", "复核")):
        for row in contract.get(arr_name) or []:
            conf = row.get("confidence")
            needs_review = (
                tier_label == "复核"
                or (isinstance(conf, (int, float)) and conf < 0.5)
                or bool(review_reason_text(row))
            )
            rows.append({
                "来源文件": source_name,
                "sha256": digest,
                "类别/原文": str(row.get("text") or row.get("value") or row.get("kind") or ""),
                "依据/方法": str(row.get("method") or ""),
                "置信度": round(float(conf), 4) if isinstance(conf, (int, float)) else "",
                "置信分层": row.get("confidence_tier") or tier_label,
                "复核原因": review_reason_text(row),
                "证据坐标": coord_of(row),
                "状态": "需人工复核" if needs_review else "候选",
            })
    return rows


def per_file_summary(status: str, reason: str, scan: dict[str, Any], contract_path: Path) -> dict[str, Any]:
    contract = scan.get("contract") or {}
    cand = contract.get("candidates") or []
    review = contract.get("review_candidates") or []
    main_candidates = len(cand)
    review_count = len(review)
    final_status = status
    if status == "ok" and review_count > 0:
        final_status = "需人工复核"
    elif status == "ok" and main_candidates == 0 and review_count == 0:
        final_status = "可交付(空)"
    elif status == "ok":
        final_status = "可交付"
    gap = reason
    if not gap:
        gaps: list[str] = []
        for row in cand + review:
            for r in row.get("review_reasons") or []:
                if r not in gaps:
                    gaps.append(r)
        if gaps:
            gap = "；".join(gaps)
    return {
        "来源文件": scan.get("source_file", ""),
        "sha256": scan.get("input_sha256", ""),
        "状态": final_status,
        "候选数": main_candidates,
        "需复核数": review_count,
        "缺口/失败原因": gap,
        "契约文件": str(contract_path),
    }


def _sheet_safe(name: str) -> str:
    return re.sub(r"[\\/?*\[\]:]", "_", name)[:31]


def write_xlsx(path: Path, ledger_rows: list[dict[str, Any]], summaries: list[dict[str, Any]],
               per_image: dict[str, list[dict[str, Any]]]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    yellow_fill = PatternFill("solid", fgColor=YELLOW)
    warn_fill = PatternFill("solid", fgColor="FCE4D6")

    # 1) 主台账：候选级明细，需复核/低置信/复核原因非空 -> 黄底
    header = ["来源文件", "sha256", "类别/原文", "依据/方法", "置信度", "置信分层", "复核原因", "证据坐标", "状态"]
    rows = [[r["来源文件"], r["sha256"][:12], r["类别/原文"], r["依据/方法"], r["置信度"],
             r["置信分层"], r["复核原因"], r["证据坐标"], r["状态"]] for r in ledger_rows]
    ws = wb.create_sheet(_sheet_safe("识图台账"))
    ws.append(header)
    for ci in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=ci)
        cell.font = head_font
        cell.fill = head_fill
    for r in rows:
        ws.append(r)
    for ri, r in enumerate(rows, start=2):
        # 黄底条件：置信分层=复核 或 置信度<0.5 或 复核原因非空
        if r[5] == "复核" or (isinstance(r[4], (int, float)) and r[4] < 0.5) or str(r[6]).strip():
            for ci in range(1, len(header) + 1):
                ws.cell(row=ri, column=ci).fill = yellow_fill
    for ci in range(1, len(header) + 1):
        letter = get_column_letter(ci)
        max_len = max([len(str(r[ci - 1])) for r in [header] + rows] or [8])
        ws.column_dimensions[letter].width = min(max_len + 3, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # 2) 图纸汇总（含失败/缺输入）
    sum_header = ["来源文件", "sha256", "状态", "候选数", "需复核数", "缺口/失败原因", "契约文件"]
    sum_rows = [[s["来源文件"], s["sha256"][:12], s["状态"], s["候选数"], s["需复核数"],
                 s["缺口/失败原因"], s["契约文件"]] for s in summaries]
    ws_sum = wb.create_sheet(_sheet_safe("图纸汇总"))
    ws_sum.append(sum_header)
    for ci in range(1, len(sum_header) + 1):
        cell = ws_sum.cell(row=1, column=ci)
        cell.font = head_font
        cell.fill = head_fill
    for r in sum_rows:
        ws_sum.append(r)
    for ri, r in enumerate(sum_rows, start=2):
        if r[2] not in ("可交付", "可交付(空)"):
            for ci in range(1, len(sum_header) + 1):
                ws_sum.cell(row=ri, column=ci).fill = warn_fill
    for ci in range(1, len(sum_header) + 1):
        letter = get_column_letter(ci)
        max_len = max([len(str(r[ci - 1])) for r in [sum_header] + sum_rows] or [8])
        ws_sum.column_dimensions[letter].width = min(max_len + 3, 60)
    ws_sum.freeze_panes = "A2"
    ws_sum.auto_filter.ref = ws_sum.dimensions

    # 3) 每张图汇总 sheet：候选人 / 复核人 / 缺口
    for fname, digested in per_image.items():
        name = str(fname)[-24:]
        ws = wb.create_sheet(_sheet_safe("汇总·" + name))
        ws.append(["指标", "值"])
        ws.cell(row=1, column=1).font = head_font
        ws.cell(row=1, column=2).font = head_font
        ws.cell(row=1, column=1).fill = head_fill
        ws.cell(row=1, column=2).fill = head_fill
        for key, value in digested:
            ws.append([key, str(value)])
        for ci in (1, 2):
            letter = get_column_letter(ci)
            max_len = max([len(str(r[ci - 1])) for r in [["指标", "值"]] + digested] or [8])
            ws.column_dimensions[letter].width = min(max_len + 4, 80)
        ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_selfcheck_md(path: Path, summaries: list[dict[str, Any]], ledger_rows: list[dict[str, Any]], total_inputs: int) -> Path:
    counts: dict[str, int] = {}
    for s in summaries:
        st = s["状态"].split("(")[0]
        counts[st] = counts.get(st, 0) + 1
    lines: list[str] = []
    lines.append("# CAD 识图批量自检汇总")
    lines.append("")
    lines.append("- 输入文件数：%d" % total_inputs)
    lines.append("- 可交付：%d" % counts.get("可交付", 0))
    lines.append("- 需人工复核：%d" % counts.get("需人工复核", 0))
    lines.append("- 失败跳过：%d" % counts.get("失败跳过", 0))
    lines.append("- 缺输入：%d" % counts.get("缺输入", 0))
    lines.append("- 全部候选行：%d（其中需人工复核 %d 行）" % (
        len(ledger_rows), sum(1 for r in ledger_rows if r["状态"] == "需人工复核")))
    lines.append("")
    lines.append("> 本项目为识图候选与证据台账，`final_quantity=false`，不构成最终工程量。")
    lines.append("")
    lines.append("## 逐图状态")
    lines.append("")
    lines.append("| 来源文件 | sha256 | 状态 | 候选数 | 需复核数 | 缺口/失败原因 | 契约文件 |")
    lines.append("|---|---|---|---|---|---|---|")
    for s in summaries:
        lines.append("| %s | `%s` | %s | %s | %s | %s | `%s` |" % (
            s["来源文件"], (s["sha256"] or "")[:16], s["状态"], s["候选数"], s["需复核数"],
            (s["缺口/失败原因"] or "").replace("|", "\\|"), s["契约文件"].replace("|", "\\|")))
    lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def rc_reason(rc: int, err: str, budget: float) -> str:
    """把 cad_scan 非零退出/超时/内存看门狗/加密/不支持版本翻译成可记台账的中文原因。"""
    err = (err or "").strip()
    low = (err + f" rc={rc}").lower()
    if rc == 124 or "timeout expired" in low or "超时" in low:
        return f"cad_scan 超时终止（预算 {budget:.0f}s）：{err[:400]}"
    if rc == -9 or ("memory" in low and ("kill" in low or "终止" in low)) or "内存闸" in low or "内存看门狗" in low:
        return f"内存看门狗终止（大图/高 CPU）：{err[:400]}"
    if "encrypt" in low or "password" in low or "加密" in low:
        return f"加密 DWG：需密码或人工解锁后再处理（{err[:400]}）"
    if "unsupported dwg version" in low or ("不支持" in low and "dwg" in low):
        return f"不支持的 DWG 版本：{err[:400]}"
    return f"cad_scan 失败（rc={rc}）：{err[:900]}"


def process_file(path: Path, outdir: Path, budget: float, cache_dir: str | None) -> dict[str, Any]:
    stem = safe_stem(path)
    contract_path = outdir / "_contract" / f"{stem}.json"
    detail_path = outdir / "_detail" / f"{stem}.json"
    json_path = outdir / "_json" / f"{stem}.json"
    digest = sha256(path)
    small = False
    try:
        size = path.stat().st_size
    except OSError:
        size = 0
    if size < TINY_MIN_BYTES:
        small = True
    tag = classify_input(path)
    if small or tag == "缺输入":
        reason = input_fail_reason(path) or "输入过小(<1KB)或为空文件"
        summary = {
            "来源文件": path.name, "sha256": digest, "状态": "缺输入",
            "候选数": 0, "需复核数": 0, "缺口/失败原因": reason, "契约文件": "",
        }
        return {"status": "缺输入", "reason": reason, "digest": digest, "summary": summary,
                "scan": None, "ledger_rows": []}
    if tag == "失败跳过":
        reason = magic_reason(path) or "魔数/头部校验失败"
        summary = {
            "来源文件": path.name, "sha256": digest, "状态": "失败跳过",
            "候选数": 0, "需复核数": 0, "缺口/失败原因": reason, "契约文件": "",
        }
        return {"status": "失败跳过", "reason": reason, "digest": digest, "summary": summary,
                "scan": None, "ledger_rows": []}
    rc, err = run_cad_scan(path, detail_path, json_path, budget, cache_dir)
    if rc != 0:
        reason = rc_reason(rc, err, budget)
        summary = {
            "来源文件": path.name, "sha256": digest, "状态": "失败跳过",
            "候选数": 0, "需复核数": 0, "缺口/失败原因": reason, "契约文件": "",
        }
        return {"status": "失败跳过", "reason": reason, "digest": digest, "summary": summary,
                "scan": None, "ledger_rows": []}
    try:
        detail = json.loads(detail_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        reason = f"详情 JSON 读取失败：{e}"
        summary = {"来源文件": path.name, "sha256": digest, "状态": "失败跳过", "候选数": 0,
                   "需复核数": 0, "缺口/失败原因": reason, "契约文件": ""}
        return {"status": "失败跳过", "reason": reason, "digest": digest, "summary": summary,
                "scan": None, "ledger_rows": []}
    scan = build_contract(detail, path.name, path, digest)
    contract_path.parent.mkdir(parents=True, exist_ok=True)
    contract_path.write_text(json.dumps(scan, ensure_ascii=False, indent=1), encoding="utf-8")
    errors = validate_contract_payload(scan)
    if errors:
        scan["_contract_errors"] = errors[:20]
    summary = per_file_summary("ok", "；".join(errors[:3]) if errors else "", scan, contract_path)
    ledger_rows = flatten_candidates(digest, scan.get("source_file", path.name), scan.get("contract") or {})
    return {"status": summary["状态"], "reason": summary["缺口/失败原因"], "digest": digest,
            "summary": summary, "scan": scan, "ledger_rows": ledger_rows}


def main(argv: list[str] | None = None) -> int:
    import pyboot
    pyboot.ensure_utf8_io()
    ap = argparse.ArgumentParser(description="按目录批量扫图并输出识图台账（Excel + 自检汇总.md）")
    ap.add_argument("paths", nargs="+", help="输入文件或目录")
    ap.add_argument("--out", default="识图台账", help="输出目录（默认 ./识图台账），写入 识图台账.xlsx、自检汇总.md、_contract/、_detail/、_json/")
    ap.add_argument("--glob", default="*.dwg", help="目录扫描扩展名（分号分隔，如 *.dwg;*.dxf；默认 *.dwg）")
    ap.add_argument("--recursive", action="store_true", help="递归扫描子目录")
    ap.add_argument("--resume", action="store_true", help="按源文件 SHA-256 断点续跑（跳过已成功文件）")
    ap.add_argument("--budget", type=float, default=45.0, help="单图 cad_scan 时间预算/秒，默认 45")
    ap.add_argument("--cache-dir", default=None, help="传给 cad_scan 的本机证据缓存目录")
    args = ap.parse_args(argv)

    patterns = [p.strip() for p in args.glob.split(";") if p.strip()] or ["*.dwg"]
    files: list[Path] = []
    for pat in patterns:
        files.extend(collect_files(args.paths, pat, args.recursive))
    files = sorted(set(files), key=lambda p: str(p))
    if not files:
        print("没有找到匹配输入文件", file=sys.stderr)
        return 2

    outdir = Path(args.out).expanduser()
    outdir.mkdir(parents=True, exist_ok=True)
    state_path = outdir / "_state.json"
    state: dict[str, dict[str, Any]] = {}
    if args.resume and state_path.exists():
        try:
            state = json.loads(state_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            state = {}

    ledger_rows: list[dict[str, Any]] = []
    summaries: list[dict[str, Any]] = []
    per_image: dict[str, list[dict[str, Any]]] = {}
    processed = 0
    resume_skipped = 0
    for f in files:
        stem = safe_stem(f)
        contract_path = outdir / "_contract" / f"{stem}.json"
        digest = sha256(f)
        if args.resume and digest in state:
            rec = state[digest]
            if rec.get("status") in ("可交付", "需人工复核", "可交付(空)") and contract_path.exists():
                try:
                    scan = json.loads(contract_path.read_text(encoding="utf-8"))
                    summary = per_file_summary("ok", "", scan, contract_path)
                    summaries.append(summary)
                    ledger_rows.extend(flatten_candidates(digest, scan.get("source_file", f.name), scan.get("contract") or {}))
                    cand = scan.get("contract", {}).get("candidates") or []
                    review = scan.get("contract", {}).get("review_candidates") or []
                    per_image[f.name] = [
                        ("来源文件", f.name), ("sha256", digest), ("状态", summary["状态"]),
                        ("候选人", len(cand)), ("复核人", len(review)),
                        ("缺口", summary["缺口/失败原因"] or "无"),
                        ("契约文件", str(contract_path)),
                    ]
                    resume_skipped += 1
                    continue
                except (OSError, json.JSONDecodeError):
                    pass
        result = process_file(f, outdir, args.budget, args.cache_dir)
        processed += 1
        summaries.append(result["summary"])
        ledger_rows.extend(result["ledger_rows"])
        if result["scan"] is not None:
            scan = result["scan"]
            contract = scan.get("contract") or {}
            cand = contract.get("candidates") or []
            review = contract.get("review_candidates") or []
            per_image[f.name] = [
                ("来源文件", f.name), ("sha256", digest), ("状态", result["summary"]["状态"]),
                ("候选人", len(cand)), ("复核人", len(review)),
                ("缺口", result["summary"]["缺口/失败原因"] or "无"),
                ("契约文件", str(contract_path)),
            ]
        state[digest] = {
            "status": result["summary"]["状态"],
            "reason": result["summary"]["缺口/失败原因"],
            "source": f.name,
            "sha256": digest,
            "contract": str(contract_path) if result["scan"] is not None else "",
            "updated": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        }
        try:
            state_path.write_text(json.dumps(state, ensure_ascii=False, indent=1), encoding="utf-8")
        except OSError:
            pass

    xlsx_path = outdir / "识图台账.xlsx"
    write_xlsx(xlsx_path, ledger_rows, summaries, per_image)
    md_path = outdir / "自检汇总.md"
    write_selfcheck_md(md_path, summaries, ledger_rows, len(files))
    result_obj = {
        "out_dir": str(outdir),
        "xlsx": str(xlsx_path),
        "selfcheck_md": str(md_path),
        "input_files": len(files),
        "processed": processed,
        "resume_skipped": resume_skipped,
        "status_counts": {
            s: sum(1 for x in summaries if x["状态"] == s or x["状态"].startswith(s))
            for s in ("可交付", "需人工复核", "失败跳过", "缺输入")
        },
    }
    print(json.dumps(result_obj, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
