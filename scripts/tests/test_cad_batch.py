#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = SKILL_DIR / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

from cad_batch import (  # noqa: E402
    build_contract,
    classify_input,
    flatten_candidates,
    sha256,
)
from cad_contract import validate_contract_payload  # noqa: E402

SAMPLE = SKILL_DIR / "examples" / "sample_room.dxf"


def make_detail(records):
    return {"text_records": records}


class CadBatchTest(unittest.TestCase):
    def test_sample_detail_builds_valid_v0_contract(self):
        detail = make_detail([
            {"kind": "TEXT", "text": "KL-1 300x600", "x": 1200.0, "y": 1000.0, "layer": "TEXT"},
            {"kind": "TEXT", "text": "C30", "x": 400.0, "y": 1500.0, "layer": "TEXT"},
        ])
        scan = build_contract(detail, "sample_room.dxf", SAMPLE, "0" * 64)
        contract = scan["contract"]
        self.assertEqual(contract["schema"], "cad-file-reader/v0")
        self.assertEqual(contract["summary"]["candidate_count"], 2)
        self.assertEqual(contract["summary"]["review_required"], 0)
        for row in contract["candidates"]:
            self.assertIs(row["final_quantity"], False)
            self.assertEqual(row["source_schema"], "cad-file-reader/v0")
        self.assertEqual(validate_contract_payload(scan), [])

    def test_low_confidence_row_is_review_and_flatten_marks_review(self):
        # 缺 bbox -> 置信 0.40 -> review_required
        detail = make_detail([
            {"kind": "TEXT", "text": "管线走向", "layer": ""},
        ])
        scan = build_contract(detail, "bad.dxf", SAMPLE, "1" * 64)
        contract = scan["contract"]
        self.assertEqual(contract["summary"]["review_required"], 1)
        rows = flatten_candidates("1" * 64, "bad.dxf", contract)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["状态"], "需人工复核")
        self.assertTrue(rows[0]["复核原因"])

    
    def test_xlsx_review_row_gets_yellow_fill(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl 不在当前解释器")
        from cad_batch import write_xlsx
        ledger = [
            {"来源文件": "ok.dxf", "sha256": "0"*64, "类别/原文": "OK", "依据/方法": "t",
             "置信度": 0.9, "置信分层": "inferred_candidate", "复核原因": "", "证据坐标": "0,0~1,1", "状态": "候选"},
            {"来源文件": "review.dxf", "sha256": "1"*64, "类别/原文": "R1", "依据/方法": "t",
             "置信度": 0.4, "置信分层": "review_required", "复核原因": "low_confidence_inference", "证据坐标": "", "状态": "需人工复核"},
        ]
        summaries = [{"来源文件": "review.dxf", "sha256": "1"*64, "状态": "需人工复核",
                      "候选数": 0, "需复核数": 1, "缺口/失败原因": "low_confidence_inference", "契约文件": ""}]
        per_image = {"review.dxf": [("候选人", "0"), ("复核人", "1"), ("缺口", "low_confidence_inference")]}
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "台账.xlsx"
            write_xlsx(out, ledger, summaries, per_image)
            wb = load_workbook(out)
            ws = wb["识图台账"]
            fills = {}
            for r in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False):
                fills[str(r[0].value)] = r[0].fill.fgColor.rgb
            self.assertIn("ok.dxf", fills)
            self.assertIn("review.dxf", fills)
            self.assertEqual(fills["review.dxf"], "00FFF2CC")
            self.assertNotEqual(fills["ok.dxf"], "00FFF2CC")

    def test_failure_and_missing_input_are_recorded_not_blocking(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            good = root / "sample_room.dxf"
            good.write_bytes(SAMPLE.read_bytes())
            bad = root / "bad.dxf"
            bad.write_bytes(b"not-a-dwg" * 200)
            tiny = root / "tiny.dxf"
            tiny.write_bytes(b"")
            outdir = root / "out"
            self.assertEqual(classify_input(tiny), "缺输入")
            self.assertEqual(classify_input(bad), "失败跳过")
            self.assertEqual(classify_input(good), "ok")
            try:
                import openpyxl  # noqa: F401
                has_openpyxl = True
            except ImportError:
                has_openpyxl = False
            if not has_openpyxl:
                self.skipTest("openpyxl 不在当前解释器")
            from cad_batch import main
            # 只跑 CLI 外层：单图/坏图不中断，能写出台账文件
            proc = subprocess.run(
                [sys.executable, str(SCRIPTS_DIR / "cad_batch.py"), str(root),
                 "--out", str(outdir), "--glob", "*.dxf"],
                capture_output=True, text=True, timeout=180,
            )
            self.assertEqual(proc.returncode, 0, proc.stderr)
            self.assertTrue((outdir / "识图台账.xlsx").exists())
            self.assertTrue((outdir / "自检汇总.md").exists())
            md = (outdir / "自检汇总.md").read_text(encoding="utf-8")
            self.assertIn("失败跳过", md)
            self.assertIn("缺输入", md)
            # 契约文件校验
            from cad_contract import validate_contract_file
            self.assertEqual(validate_contract_file(outdir / "_contract" / "sample_room.json"), [])


if __name__ == "__main__":
    unittest.main()
