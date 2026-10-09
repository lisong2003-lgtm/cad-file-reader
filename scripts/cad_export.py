#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""cad_export.py — 把 cad_scan 结构化结果导出为 Excel(.xlsx) / Word(.docx)。

用法（内部函数，供 cad_scan.py 调用）：
    from cad_export import write_xlsx, write_docx
    write_xlsx(path, sections)
    write_docx(path, title, sections)
其中 sections 为 [(sheet/章节名, 表头列表, 行二维列表), ...]
仅输出候选/中间数据（final_quantity=false），不产生工程量。
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, Sequence


def _norm_cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


def write_xlsx(path: str | Path, sections: Sequence[tuple[str, Sequence[str], Iterable[Sequence[Any]]]]) -> Path:
    """写出多 Sheet 的 xlsx；sections 每项为 (sheet名, 表头, 行列表)。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    for sheet_name, header, rows in sections:
        ws = wb.create_sheet(sheet_name[:31])
        if header:
            ws.append([str(_norm_cell(c)) for c in header])
            for col in range(1, len(header) + 1):
                cell = ws.cell(row=1, column=col)
                cell.font = head_font
                cell.fill = head_fill
        for row in rows:
            ws.append([_norm_cell(c) for c in row])
        if header:
            for col in range(1, len(header) + 1):
                letter = get_column_letter(col)
                max_len = 8
                for row in ws.iter_rows(min_col=col, max_col=col, min_row=1):
                    for cell in row:
                        max_len = max(max_len, len(str(cell.value or "")))
                ws.column_dimensions[letter].width = min(max_len + 2, 80)
        ws.freeze_panes = "A2"
    wb.save(p)
    return p


def write_docx(path: str | Path, title: str, sections: Sequence[tuple[str, Sequence[str], Iterable[Sequence[Any]]]]) -> Path:
    """写出分节 Word 文档；sections 每项为 (章节名, 表头, 行列表)。"""
    from docx import Document
    from docx.shared import Pt

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    h = doc.add_heading(title, level=0)
    for run in h.runs:
        run.font.size = Pt(16)
    para = doc.add_paragraph()
    para.add_run("说明：本文件仅为 CAD 识图候选/证据，final_quantity=false，不替代专业算量。")
    for section_name, header, rows in sections:
        doc.add_heading(section_name, level=1)
        if header:
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for i, c in enumerate(header):
                table.rows[0].cells[i].text = str(c)
        else:
            table = doc.add_table(rows=1, cols=1)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = ""
        for row in rows:
            cells = table.add_row().cells
            for i in range(min(len(cells), max(len(row), 1))):
                cells[i].text = str(_norm_cell(row[i] if i < len(row) else ""))
        doc.add_paragraph()
    doc.save(p)
    return p
